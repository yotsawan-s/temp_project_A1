from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

NAVY="1F3864"; BLUE="2E5496"; LTBLUE="D6E0F0"; GREY="F2F2F2"
YELLOW="FFF2CC"; INPUT="FFFF00"; GREEN="C6EFCE"; GREENT="006100"
RED="FFC7CE"; REDT="9C0006"; ORANGE="FFEB9C"; ORANGET="9C5700"; WHITE="FFFFFF"
F="Arial"
thin=Side(style="thin",color="BFBFBF")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

def st(c,*,b=False,sz=10,color="000000",fill=None,wrap=False,h="left",v="center",bd=False):
    c.font=Font(name=F,bold=b,size=sz,color=color)
    if fill: c.fill=PatternFill("solid",fgColor=fill)
    c.alignment=Alignment(horizontal=h,vertical=v,wrap_text=wrap)
    if bd: c.border=border

def title(ws,text,sub,ncol):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncol)
    st(ws.cell(1,1,text),b=True,sz=15,color=WHITE,fill=NAVY); ws.row_dimensions[1].height=26
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ncol)
    st(ws.cell(2,1,sub),sz=10,color=WHITE,fill=BLUE,wrap=True); ws.row_dimensions[2].height=34

def hdr(ws,row,headers,start=1,fill=BLUE,color=WHITE):
    for i,h in enumerate(headers):
        st(ws.cell(row,start+i,h),b=True,sz=10,color=color,fill=fill,h="center",wrap=True,bd=True)
    ws.row_dimensions[row].height=30

# ---- shared config ----
keymap={'3_O001':'FI Arrangement Number','1_A007':'FI Arrangement Number','1_A005':'FI Arrangement Number',
'1_A001':'FI Arrangement Number','1_A002':'Arrangement Number','3_O003':'CompositeOptSeqNum','3_O002':'OptionNum',
'1_A008':'Reference Transaction Number','4_C001':'Transaction No + DR/CR','4_C002':'Ccy + Notional','4_C003':'Currency + Accrued Int',
'2_D001':'TRADE_REF','1_A009':'FI Arrangement Number + CMF CODE','2_D002':'FI Arrangement Number','1_A003':'FI Arrangement Number',
'1_A004':'Arr Number','1_A006':'FI Arrangement Number','2_D004':'(none)','2_D005':'(none)',
'1_B001':'Arrangement Number','1_B002':'Fi Arrangement Number','1_B003':'Fi Arrangement Number','1_B004':'Arrangement Number',
'1_B005':'Fi Arrangement Number','1_B006':'Fi Arrangement Number','1_B007':'Fi Arrangement Number','1_B008':'Ref.No.',
'1_B009':'Fi Arrangement Number','2_D003':'Ref + CCY + Amount + CMF Code'}
datemap={'3_O001':'Arrangement Contract Date','1_A007':'Data Set Date','1_A005':'Data Set Date','1_A001':'Data Set Date',
'1_A002':'Data Set Date','3_O003':'Date','3_O002':'(none)','1_A008':'Data Set Date','4_C001':'(none)','4_C002':'Settlement Date',
'4_C003':'Settlement Date','2_D001':'TRADE_DATE','1_A009':'Data Set Date','2_D002':'Data Set Date','1_A003':'Data Set Date',
'1_A004':'Data Set Date','1_A006':'Data Set Date','2_D004':'(none)','2_D005':'(none)','2_D003':'Date'}
for b in ['1_B001','1_B002','1_B003','1_B004','1_B005','1_B006','1_B007','1_B008','1_B009']: datemap[b]='(none)'

# Col.Filter (Y) per source — used in Step1 Filter (from 3_DataSource)
filtermap={'2_D001':['PRODUCT_TYPE','TRADE_DATE'],'2_D002':['Data Set Date'],
'2_D003':['Date','CP Name','Type','Settle/Unsettle'],'2_D004':[],'2_D005':[],
'3_O001':['Arrangement Contract Date'],'3_O002':['Expired/Exercised'],'3_O003':['CompositeOptSeqNum','Date'],
'4_C001':[],'4_C002':['Settlement Date','Class','Ccy'],'4_C003':['Settlement Date','Status'],
'1_A001':['DEPT CODE','System ID','Data Set Date'],'1_A002':['DEPT CODE','Data Set Date'],
'1_A003':['Data Set Date'],'1_A004':['Dept Code','Data Set Date'],'1_A005':['DEPT CODE','System Id','Data Set Date'],
'1_A006':['Dept Code','Data Set Date'],'1_A007':['Dept Code','System Id','Data Set Date'],
'1_A008':['Dept Code','System Id','Data Set Date'],'1_A009':['Data Set Date']}
for b in ['1_B001','1_B002','1_B003','1_B004','1_B005','1_B006','1_B007','1_B008','1_B009']: filtermap[b]=['Department Code']

# (product,cat,set_no,chain,transmap,pattern,matchkey,compare,condition,backcheck)
P1,P2,P3="1 Direct","2 Direct+Cond","3 Map+Back-check"
sets=[
("D","D1",1,["2_D001","1_A009","1_B009"],"1 : 1 : n",P2,"Ref (TRADE_REF=FI Arr No)","CCY + Amount + Date","PRODUCT_TYPE in scope",""),
("D","D1",2,["2_D001","1_A005","1_B002"],"1 : 1 : n",P1,"Ref (FI Arr No)","CCY + Buy/Sell Amount + Date","",""),
("D","D1",3,["2_D001","1_A001","1_B003"],"1 : 1,2 : n",P2,"Ref (FI Arr No)","CCY + Original Amount + Date","Leg Type (1,2)",""),
("D","D1",4,["2_D001","1_A002","1_B004"],"1 : 1 : n",P1,"Ref (Arr Number)","CCY + Original Amount + Date","",""),
("D","D2",5,["2_D002","1_A003","2_D001","1_B006"],"1 : 1 : 1 : n",P3,"Ref (FI Arr No)","CCY + Amount + Date","","2_D001 (Main Source)"),
("D","D3",6,["2_D003","1_A001","1_B003"],"1 : 1 : n",P2,"Ref + CCY + Amount","CCY + Amount + Date","Type / Settle status",""),
("D","D3",7,["2_D003","1_A002","1_B004"],"1 : 1 : n",P2,"Ref + CCY + Amount","CCY + Amount + Date","Type / Settle status",""),
("D","D3",8,["2_D003","1_A004","1_B001"],"1 : 1 : n",P2,"Ref + CCY + Amount","CCY + Amount + Date","Type / Settle status",""),
("D","D3",9,["2_D003","1_A005","1_B002"],"1 : 1 : n",P2,"Ref + CCY + Amount","CCY + Amount + Date","Type / Settle status",""),
("D","D3",10,["2_D003","1_A008","1_B008"],"1 : 1 : n",P2,"Ref + CCY + Amount","CCY + Amount + Date","Transaction Type",""),
("D","D3",11,["2_D003","1_A006","1_B005"],"1 : n : n",P2,"Ref + Type + CCY + Amount","CCY + Amount + Date","Type",""),
("D","D3",12,["2_D003","2_D004","2_D005","1_A005"],"n : 1 : 1 : ?",P3,"Ref (CMF/Trans Ref)","CCY + Current Principal","","2_D004 + 2_D005"),
("O","O1",1,["3_O001","1_A007","1_B007"],"1 : 1 : n",P1,"Ref (FI Arr No)","CCY + Premium Amount + Date","",""),
("O","O1",2,["3_O001","1_A005","1_B002"],"1 : 1 : n",P1,"Ref (FI Arr No)","CCY + Buy/Sell Amount + Date","",""),
("O","O1",3,["3_O001","1_A001","3_O003","1_B003"],"1 : 1 : 1 : n",P3,"Ref (FI Arr No / OptSeqNum)","CCY + Amount + Premium","","3_O003 (back-check premium)"),
("O","O1",4,["3_O001","1_A002","1_B004"],"1 : 1 : n",P1,"Ref (Arr Number)","CCY + Original Amount + Date","",""),
("O","O2",5,["3_O002","1_A005","1_B002"],"1 : 2 : n",P2,"Ref (OptionNum)","CCY + Amount","Expired/Exercised status",""),
("O","O3",6,["3_O003","1_A008","1_B008"],"1 : 1 : n",P1,"Ref (OptSeqNum)","Premium CCY + Premium Amount + Date","",""),
("C","C1",1,["1_A008","4_C001"],"1 : 2",P2,"Ref (Transaction No) + DR/CR","Amount","DR/CR split",""),
("C","C1",2,["1_A008","4_C002","4_C003"],"1 : 1 : 1",P3,"Ref + Ccy + Notional","Ccy + Notional + Accrued Int","","4_C003 (back-check interest)"),
("C","C1",3,["1_A008","1_B008"],"1 : n",P1,"Ref (Ref.No.)","CCY + Amount","",""),
]

wb=Workbook()
# ============================================================ README
ws=wb.active; ws.title="README"
title(ws,"PART RECONCILE  —  BOT Reconciliation Tool  (Prototype v0.2)",
      "ตรวจ Transaction ของ A1_System (1_A / 1_B) เทียบ Product Database (D/O/C) ครบถ้วน-ถูกต้อง ตาม Data Transaction Date  |  v0.2: ครบ 21 Sets + Key Prep + 3 Match Patterns + เทียบ CCY/Amount/Date",6)
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=34
for c in "CDEF": ws.column_dimensions[c].width=22
r=4
def line(text,**kw):
    global r
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=6)
    st(ws.cell(r,2,text),**kw); r+=1
line("1)  การ Reconcile 2 มุมมอง (ต่อ 1 Match Key)",b=True,sz=12,color=NAVY)
line("    • มุม1  A1→DB : มีใน A1 แต่ไม่มีใน Database = EXTRA in A1 (บันทึกเกิน)")
line("    • มุม2  DB→A1 : มีใน Database แต่ไม่มีใน A1 = MISSING in A1 (บันทึกขาด)")
line("    • Key ตรงกันทั้งคู่ → เทียบ Compare Fields (CCY + Amount + Date) ถ้าต่าง = AMOUNT/VALUE MISMATCH")
line("    • ไม่รวมยอด (no SUM) — จับคู่รายบรรทัดด้วย Match Key | ผลลัพธ์ ERROR ผูกกับ Key → นำกลับไปแก้ที่ A1_System (ไฟล์ 1_Bxxx)")
r+=1
line("1.5)  กระบวนการ Mapping มาตรฐาน 4 Step (ต่อคู่ Main ⇄ Map File — ดูชีท Mapping_Steps)",b=True,sz=12,color=NAVY)
line("    Step1 Filter (กรองแถว) → Step2 Key Map (สร้างคีย์ Ref+อื่นๆ) → Step3 Reconcile (เทียบ Date/CCY/Amount) → Step4 Result(Main)/Return(Map→1_B)")
line("    แต่ละ Set = Main (Source#1) เทียบกับ Map File #1..#n ทีละคู่")
r+=1
line("2)  Match Patterns 3 แบบ (กำหนดต่อ Set ในชีท Recon_Rules)",b=True,sz=12,color=NAVY)
line("    แบบ 1  Direct map         : จับคู่ตรงด้วย Match Key แล้วเทียบ Compare Fields")
line("    แบบ 2  Direct + Condition : กรองเงื่อนไขก่อน (เช่น Product/Transaction Type, Settle status) แล้วจึงจับคู่")
line("    แบบ 3  Map + Back-check   : จับคู่ผ่านสายโซ่ แล้วย้อนกลับไปตรวจ Main Source (สำหรับ chain 3-4 ขั้น)")
r+=1
line("3)  Key Prep (เตรียมคีย์ก่อน mapping) — DSL ในชีท Recon_Rules",b=True,sz=12,color=NAVY)
line('    AS_IS  |  LEFT(n) | RIGHT(n) | MID(start,len) | AFTER("x") | BEFORE("x") | REGEX("pattern") | CONCAT(f1,f2)')
line('    ตัวอย่าง: คีย์ดิบ "KT20260530C0001"  →  REGEX("C\\d+")  →  "C0001"  (ใช้จับคู่กับ Database)')
r+=1
line("4)  Match Key forms (ฝั่ง n ไม่รวมยอด)",b=True,sz=12,color=NAVY)
line("    Ref อย่างเดียว  |  Ref + CCY + Amount  |  Ref + Type + CCY + Amount   (เลือกต่อ Set)")
r+=1
line("4.5)  Normalize ค่า Date & Number ก่อนเทียบ (ชีท Field_Format)",b=True,sz=12,color=NAVY)
line("    • Date: แต่ละ Source อาจเก็บคนละ format (DATE / NUMBER:yyyymmdd / TEXT:dd/mm/yyyy / TEXT:dd-mmm-yy ฯลฯ)")
line("            → Engine แปลงเป็น canonical ISO (yyyy-mm-dd) ก่อนเทียบ, Tool แสดงผลเป็น dd/mm/yyyy")
line('    • Number: ตัดเครื่องหมาย/แปลง sign ก่อนเทียบ เช่น -100 → ABS=100 , หรือ SIGN_BY("DR/CR") , "(100)"→-100')
r+=1
line("5)  วิธีใช้",b=True,sz=12,color=NAVY)
line("    Control_Panel (กรอกเหลือง) → Recon_Rules (กำหนด logic ต่อ Set) → Sample_Data (จำลอง) → Recon_Detail (ผลรายบรรทัด) → Recon_Summary (Dashboard 21 Sets)")
r+=1
line("6)  คำอธิบายสี (Legend)",b=True,sz=12,color=NAVY)
for txt,fl,co in [("ช่องกรอก/แก้ไข (Input/Config)",INPUT,"000000"),("OK / ตรงกัน",GREEN,GREENT),
                  ("MISSING in A1 (บันทึกขาด)",RED,REDT),("EXTRA in A1 (บันทึกเกิน)",RED,REDT),
                  ("DUPLICATE KEY (คีย์ซ้ำ — ผิด Trans.Map 1:1)",RED,REDT),("VALUE MISMATCH (ค่าไม่ตรง)",ORANGE,ORANGET)]:
    st(ws.cell(r,2,"  "),fill=fl,bd=True); st(ws.cell(r,3,txt),color=co,b=True); r+=1
ws.sheet_view.showGridLines=False

# ============================================================ CONTROL PANEL
ws=wb.create_sheet("Control_Panel")
title(ws,"CONTROL PANEL  —  พารามิเตอร์การรัน",
      "กรอกช่องสีเหลือง — ใช้เป็นตัวกรอง Col.Filter (วันที่/หน่วยงาน) ของทุกไฟล์",4)
ws.sheet_view.showGridLines=False
for c,w in zip("ABCD",[3,30,26,42]): ws.column_dimensions[c].width=w
inputs=[("Data Transaction Date","2026-05-30","วันที่ Transaction ที่ตรวจ (Data Set Date/TRADE_DATE/Settlement Date/Date)"),
        ("Department Code","0123","รหัสหน่วยงาน (Dept Code/DEPT CODE)"),
        ("Product to run","ALL","D / O / C / ALL"),
        ("Amount tolerance","0","ผลต่างจำนวนเงินที่ยอมรับ (0 = ตรงเป๊ะ)")]
st(ws.cell(4,2,"พารามิเตอร์"),b=True,color=WHITE,fill=BLUE,bd=True)
st(ws.cell(4,3,"ค่า (กรอก)"),b=True,color=WHITE,fill=BLUE,bd=True,h="center")
st(ws.cell(4,4,"คำอธิบาย"),b=True,color=WHITE,fill=BLUE,bd=True)
r=5; names={}
for lab,val,desc in inputs:
    st(ws.cell(r,2,lab),b=True,bd=True); st(ws.cell(r,3,val),fill=INPUT,bd=True,h="center")
    st(ws.cell(r,4,desc),wrap=True,bd=True,sz=9); names[lab]=f"Control_Panel!$C${r}"; r+=1
wb.defined_names.add(DefinedName("ChkDate",attr_text=names["Data Transaction Date"]))
wb.defined_names.add(DefinedName("Tol",attr_text=names["Amount tolerance"]))
r+=1
# dropdown: Product to run (row 7 = 3rd input)
dv=DataValidation(type="list",formula1='"D,O,C,ALL"',allow_blank=False)
ws.add_data_validation(dv); dv.add(ws.cell(7,3))
st(ws.cell(r,2,"สถานะรวม (Overall KPI — อ้างอิง Recon_Summary)"),b=True,sz=12,color=NAVY); r+=1
SUMST=f"Recon_Summary!$K$5:$K${4+len(sets)}"
for lab,fm,fl in [("Set ทั้งหมด",f'=COUNTA(Recon_Summary!$A$5:$A${4+len(sets)})',GREY),
                  ("Set ผ่าน (OK)",f'=COUNTIF({SUMST},"OK")',GREEN),
                  ("Set พบ Error",f'=COUNTIF({SUMST},"ERROR*")',RED),
                  ("Set รอข้อมูล (Pending)",f'=COUNTIF({SUMST},"Pending*")',YELLOW)]:
    st(ws.cell(r,2,lab),b=True,bd=True)
    ws.cell(r,3,fm); st(ws.cell(r,3),bd=True,h="center",b=True,fill=fl); r+=1

# ============================================================ RECON SUMMARY
ws=wb.create_sheet("Recon_Summary")
title(ws,"RECON SUMMARY  —  Dashboard สรุปทุก Set (21 Sets)",
      "หนึ่งแถว/Set | สีแดง = พบ Error ต้องแก้ | Pattern = วิธีจับคู่ | Set D1-1 คำนวณสดจาก Sample_Data, Set อื่น = Pending (รอโหลดข้อมูล)",10)
ws.sheet_view.showGridLines=False
hdr(ws,4,["Product","Cat","Set","Chain (Database ⇄ A1_System)","Trans.Map","Pattern","MISSING in A1","EXTRA in A1","VALUE MISMATCH","DUP KEY","Overall Status"])
ws.freeze_panes="A5"
DET_LAST=11  # Recon_Detail data rows 6..11
r=5
for p,cat,no,chain,tm,pat,mk,cmp,cond,bc in sets:
    for i,v in enumerate([p,cat,no," ⇄ ".join(chain),tm,pat]):
        st(ws.cell(r,i+1,v),bd=True,sz=9,h="center" if i in(0,1,2,4,5) else "left")
    if (p,cat,no)==("D","D1",1):
        ws.cell(r,7,f'=COUNTIF(Recon_Detail!$K$6:$K${DET_LAST},"MISSING in A1")')
        ws.cell(r,8,f'=COUNTIF(Recon_Detail!$K$6:$K${DET_LAST},"EXTRA in A1")')
        ws.cell(r,9,f'=COUNTIF(Recon_Detail!$K$6:$K${DET_LAST},"VALUE MISMATCH")')
        ws.cell(r,10,f'=COUNTIF(Recon_Detail!$K$6:$K${DET_LAST},"DUPLICATE*")')
        ws.cell(r,11,f'=IF(SUM(G{r}:J{r})=0,"OK","ERROR — "&SUM(G{r}:J{r})&" รายการ")')
        for cc in range(7,12): st(ws.cell(r,cc),bd=True,h="center",b=(cc==11))
    else:
        for cc in range(7,11): st(ws.cell(r,cc,"-"),bd=True,h="center")
        st(ws.cell(r,11,"Pending data"),bd=True,h="center",color="808080")
    r+=1
slast=r-1
for col,w in zip("ABCDEFGHIJK",[8,6,5,38,13,15,13,12,15,10,18]): ws.column_dimensions[col].width=w
ws.conditional_formatting.add(f"K5:K{slast}",FormulaRule(formula=['$K5="OK"'],fill=PatternFill("solid",fgColor=GREEN),font=Font(name=F,color=GREENT,bold=True)))
ws.conditional_formatting.add(f"K5:K{slast}",FormulaRule(formula=['LEFT($K5,5)="ERROR"'],fill=PatternFill("solid",fgColor=RED),font=Font(name=F,color=REDT,bold=True)))

# ============================================================ RECON RULES (engine spec)
ws=wb.create_sheet("Recon_Rules")
title(ws,"RECON RULES  —  สเปกการจับคู่ต่อ Set (อ่านได้โดย Engine / AI / คน)",
      "หัวใจของเครื่องมือ: กำหนดต่อ Set ว่าใช้ Pattern ใด, Match Key อะไร, เตรียมคีย์ (Key Prep) อย่างไร, เทียบ field ใด, เงื่อนไข/Back-check  |  ช่องเหลือง = แก้ไขได้ (ค่า default โปรดยืนยัน)",13)
ws.sheet_view.showGridLines=False
hdr(ws,4,["Set ID","Chain","Trans.Map","Match Pattern","Match Key Fields","DB Key (raw)","DB Key Prep",
          "A1 Key (raw)","A1 Key Prep","Compare Fields","Extra Condition (P2)","Back-check Source (P3)","Notes"])
ws.freeze_panes="D5"
r=5
for p,cat,no,chain,tm,pat,mk,cmp,cond,bc in sets:
    sid=f"{cat}-{no}"
    dbk=keymap.get(chain[0],"")
    a1=next((x for x in chain if x.startswith("1_A")),"")
    a1k=keymap.get(a1,a1 or "-")
    dbprep="AS_IS"; a1prep="AS_IS"
    note=""
    if sid=="D1-1": a1prep='REGEX("C\\d+")  (เช่น KT…C0001→C0001)'; note="ตัวอย่าง live ใน Recon_Detail"
    row=[sid," ⇄ ".join(chain),tm,pat,mk,dbk,dbprep,a1k,a1prep,cmp,cond if pat==P2 else "-",bc if pat==P3 else "-",note]
    for i,v in enumerate(row):
        st(ws.cell(r,i+1,v),bd=True,sz=9,h="center" if i in(2,3) else "left")
    st(ws.cell(r,1),b=True,fill=GREY,h="center")
    for col in (4,5,7,9,10,11,12): ws.cell(r,col).fill=PatternFill("solid",fgColor=INPUT)  # editable config
    r+=1
rlast=r-1
for i,w in enumerate([8,34,13,16,24,22,22,22,26,28,22,24,26]): ws.column_dimensions[get_column_letter(i+1)].width=w
# DSL note
r+=1
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=13)
st(ws.cell(r,1,'Key Prep DSL:  AS_IS | LEFT(n) | RIGHT(n) | MID(start,len) | AFTER("x") | BEFORE("x") | REGEX("pattern") | CONCAT(f1,f2)        '
                'Pattern: 1 Direct / 2 Direct+Cond / 3 Map+Back-check'),b=True,sz=9,color=NAVY,fill=YELLOW,wrap=True)
ws.row_dimensions[r].height=28

# ============================================================ MAPPING STEPS (4-step process)
ws=wb.create_sheet("Mapping_Steps")
title(ws,"MAPPING STEPS  —  กระบวนการ Mapping มาตรฐาน 4 Step (ต่อคู่ Main File ⇄ Map File)",
      "แต่ละ Set ทำทีละคู่: Main (Source#1) เทียบกับ Map File #1..#n  |  Step1 Filter → Step2 Key Map → Step3 Reconcile → Step4 Result(Main)/Return(Map)  |  ไฟล์ 1_Bxxx = ที่เขียนผล Return",8)
ws.sheet_view.showGridLines=False
STP="DDEBF7"
hdr(ws,4,["Set ID","Pairing  (Main ⇄ Map #k)","Step1 Filter — Main","Step1 Filter — Map",
          "Step2 Key Map — Main","Step2 Key Map — Map","Step3 Reconcile (หลัง Normalize)","Step4 Result → Return (Output)"])
ws.freeze_panes="C5"
r=5
for p,cat,no,chain,tm,pat,mk,cmp,cond,bc in sets:
    sid=f"{cat}-{no}"
    main=chain[0]
    outs=[x for x in chain if x.startswith("1_B")]
    out=outs[0] if outs else "(compare only)"
    maps=[x for x in chain[1:] if not x.startswith("1_B")]
    for k,mp in enumerate(maps,1):
        f_main=", ".join(filtermap.get(main,[])) or "-"
        f_map=", ".join(filtermap.get(mp,[])) or "-"
        recon="Date + CCY + Amount + อื่นๆ"
        result=f"Result(Main) / Return → {out}"
        row=[sid,f"{main} ⇄ {mp}  (#{k})",f_main,f_map,keymap.get(main,""),keymap.get(mp,""),recon,result]
        for i,v in enumerate(row): st(ws.cell(r,i+1,v),bd=True,sz=9,h="center" if i==0 else "left")
        st(ws.cell(r,1),b=True,fill=GREY,h="center"); st(ws.cell(r,2),b=True)
        for col in (3,4): ws.cell(r,col).fill=PatternFill("solid",fgColor=STP)
        for col in (5,6): ws.cell(r,col).fill=PatternFill("solid",fgColor="FCE4D6")
        st(ws.cell(r,7),fill="E2EFDA"); st(ws.cell(r,8),fill=GREEN,sz=9,bd=True)
        r+=1
mlast=r-1
for i,w in enumerate([9,26,26,28,22,22,24,30]): ws.column_dimensions[get_column_letter(i+1)].width=w
r+=1
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)
st(ws.cell(r,1,'Step1 Filter = กรองแถวในขอบเขต (วันที่/Dept/System/Type/Status)  •  Step2 Key Map = สร้างคีย์ (Ref + อื่นๆ) พร้อม Key Prep  •  '
              'Step3 Reconcile = เทียบ Date+CCY+Amount(+อื่นๆ) หลัง Normalize (ดู Field_Format)  •  '
              'Step4 = ฝั่ง Main ได้ "Result: Reconcile", ฝั่ง Map เขียน "Return: Reconcile" (ERROR Message) กลับไฟล์ 1_Bxxx'),
   b=True,sz=9,color=NAVY,fill=YELLOW,wrap=True); ws.row_dimensions[r].height=46

# ============================================================ SAMPLE DATA
ws=wb.create_sheet("Sample_Data")
title(ws,"SAMPLE DATA  —  Set D1-1 (2_D001 ⇄ 1_A009): Key Prep + Date/Number Normalization",
      "DB เก็บวันที่เป็น NUMBER(yyyymmdd) | A1 เก็บเป็น TEXT(dd/mm/yyyy) | จำนวนเงินมีเครื่องหมายลบต่างกัน → ต้อง Normalize ก่อนเทียบ | คอลัมน์เขียว = ค่าหลัง Normalize (canonical ISO / ABS)",13)
ws.sheet_view.showGridLines=False
# DB table : date stored as NUMBER yyyymmdd, amount may carry sign
st(ws.cell(4,2,"ตาราง A : Database 2_D001  (TRADE_DATE = NUMBER yyyymmdd)"),b=True,color=WHITE,fill="538135"); ws.merge_cells("B4:G4")
hdr(ws,5,["TRADE_REF (Key)","CCY","Amount (raw)","TRADE_DATE (raw)","Norm Amt (ABS)","Norm Date (ISO)"],start=2,fill="A9D08E",color="000000")
# (key, ccy, amount_raw, date_raw_num)
db_rows=[("C0001","USD",1000000,20260530),("C0002","USD",2000000,20260530),
         ("C0003","EUR",1500000,20260530),("C0004","JPY",3000000,20260530),
         ("C0006","USD",500000,20260530)]
r=6
for k,c,a,d in db_rows:
    st(ws.cell(r,2,k),bd=True); st(ws.cell(r,3,c),bd=True,h="center")
    cc=ws.cell(r,4,a); st(cc,bd=True,h="right"); cc.number_format="#,##0"
    st(ws.cell(r,5,d),bd=True,h="center")
    ws.cell(r,6,f'=ABS(D{r})'); st(ws.cell(r,6),bd=True,h="right",fill=GREEN,b=True); ws.cell(r,6).number_format="#,##0"
    ws.cell(r,7,f'=TEXT(DATE(INT(E{r}/10000),INT(MOD(E{r},10000)/100),MOD(E{r},100)),"yyyy-mm-dd")')
    st(ws.cell(r,7),bd=True,h="center",fill=GREEN,b=True)
    r+=1
db_last=r-1
# A1 table : raw key needs prep, date stored as TEXT dd/mm/yyyy, amount sign differs
st(ws.cell(4,9,"ตาราง B : A1_System 1_A009  (Data Set Date = TEXT dd/mm/yyyy) — คอลัมน์เขียว = Prep/Normalize"),b=True,color=WHITE,fill="2E75B6"); ws.merge_cells("I4:O4")
hdr(ws,5,["Raw Key","CCY","Amount (raw)","Date (raw text)","Prepped Key","Norm Amt (ABS)","Norm Date (ISO)"],start=9,fill="9DC3E6",color="000000")
# (raw_key, ccy, amount_raw, date_raw_text)  — note signs & formats differ from DB
a1_rows=[("KT20260530C0001","USD",-1000000,"30/05/2026"),  # sign- & date-format differ -> still OK after normalize
         ("KT20260530C0002","USD", 2050000,"30/05/2026"),  # amount truly differs -> MISMATCH
         ("KT20260530C0003","EUR", 1500000,"29/05/2026"),  # date truly differs -> MISMATCH
         ("KT20260530C0005","GBP",  800000,"30/05/2026"),  # extra
         ("KT20260530C0006","USD",  500000,"30/05/2026"),  # duplicate key in A1 (recorded twice)
         ("KT20260530C0006","USD",  500000,"30/05/2026")]
r=6
for k,c,a,d in a1_rows:
    st(ws.cell(r,9,k),bd=True); st(ws.cell(r,10,c),bd=True,h="center")
    cc=ws.cell(r,11,a); st(cc,bd=True,h="right"); cc.number_format="#,##0;(#,##0)"
    st(ws.cell(r,12,d),bd=True,h="center")
    ws.cell(r,13,f'=MID(I{r},FIND("C",I{r}),100)'); st(ws.cell(r,13),bd=True,h="center",fill=GREEN,b=True)          # Key Prep
    ws.cell(r,14,f'=ABS(K{r})'); st(ws.cell(r,14),bd=True,h="right",fill=GREEN,b=True); ws.cell(r,14).number_format="#,##0"  # ABS
    ws.cell(r,15,f'=TEXT(DATE(VALUE(RIGHT(L{r},4)),VALUE(MID(L{r},4,2)),VALUE(LEFT(L{r},2))),"yyyy-mm-dd")')         # parse text date
    st(ws.cell(r,15),bd=True,h="center",fill=GREEN,b=True)
    r+=1
a1_last=r-1
note=('Normalize ก่อนเทียบ:  C0001 = OK (DB amt 1,000,000 vs A1 -1,000,000 → ABS ตรง; วันที่คนละ format → ISO ตรง)  |  '
      'C0002 = MISMATCH(Amount จริง)  |  C0003 = MISMATCH(Date จริง 29 vs 30)  |  C0004 = MISSING in A1  |  C0005 = EXTRA in A1  |  '
      'C0006 = DUPLICATE in A1 (บันทึกซ้ำ 2 ครั้ง — ผิด Trans.Map 1:1)')
ws.merge_cells("B13:O14"); st(ws.cell(13,2,note),wrap=True,color="9C5700",fill=YELLOW,bd=True)
for col,w in zip("BCDEFGHIJKLMNO",[15,7,13,14,13,13,3,16,7,13,13,11,13,13]): ws.column_dimensions[col].width=w
# canonical ranges used by Recon_Detail (compare on NORMALIZED columns)
DBK=f"Sample_Data!$B$6:$B${db_last}"; DBC=f"Sample_Data!$C$6:$C${db_last}"; DBA=f"Sample_Data!$F$6:$F${db_last}"; DBD=f"Sample_Data!$G$6:$G${db_last}"
A1K=f"Sample_Data!$M$6:$M${a1_last}"; A1C=f"Sample_Data!$J$6:$J${a1_last}"; A1A=f"Sample_Data!$N$6:$N${a1_last}"; A1D=f"Sample_Data!$O$6:$O${a1_last}"

# ============================================================ RECON DETAIL
ws=wb.create_sheet("Recon_Detail")
title(ws,"RECON DETAIL  —  Set D1-1 (Pattern 1) | จับคู่ Prepped Key, เทียบค่าหลัง Normalize (CCY + Amount(ABS) + Date(ISO))",
      "ค่า Amount/Date ที่แสดง = หลัง Normalize แล้ว (Amount=ABS, Date=dd/mm/yyyy มาตรฐาน Tool) | Field Diff บอก field ที่ไม่ตรงจริง | ERROR Message → นำกลับไปแก้ที่ A1_System (ไฟล์ 1_B009)",12)
ws.sheet_view.showGridLines=False
hdr(ws,5,["Match Key","In DB?","In A1?","DB CCY","A1 CCY","DB Amount (norm)","A1 Amount (norm)","DB Date (dd/mm/yyyy)","A1 Date (dd/mm/yyyy)","Field Diff","Result / Status","ERROR Message (→ fix in A1_System)"])
ws.freeze_panes="B6"
union=["C0001","C0002","C0003","C0004","C0005","C0006"]
r=6
for k in union:
    R=r
    st(ws.cell(R,1,k),bd=True,b=True)
    ws.cell(R,2,f'=IF(COUNTIF({DBK},A{R})>0,"YES","NO")');
    ws.cell(R,3,f'=IF(COUNTIF({A1K},A{R})>0,"YES","NO")')
    ws.cell(R,4,f'=IFERROR(INDEX({DBC},MATCH(A{R},{DBK},0)),"")')
    ws.cell(R,5,f'=IFERROR(INDEX({A1C},MATCH(A{R},{A1K},0)),"")')
    ws.cell(R,6,f'=IFERROR(INDEX({DBA},MATCH(A{R},{DBK},0)),"")')
    ws.cell(R,7,f'=IFERROR(INDEX({A1A},MATCH(A{R},{A1K},0)),"")')
    ws.cell(R,8,f'=IFERROR(TEXT(DATEVALUE(INDEX({DBD},MATCH(A{R},{DBK},0))),"dd/mm/yyyy"),"")')
    ws.cell(R,9,f'=IFERROR(TEXT(DATEVALUE(INDEX({A1D},MATCH(A{R},{A1K},0))),"dd/mm/yyyy"),"")')
    # field diff only meaningful when both present
    ws.cell(R,10,(f'=IF(AND(B{R}="YES",C{R}="YES"),TRIM(IF(D{R}<>E{R},"CCY ","")&IF(ABS(F{R}-G{R})>Tol,"Amount ","")&IF(H{R}<>I{R},"Date ","")),"")'))
    ws.cell(R,11,(f'=IF(COUNTIF({A1K},A{R})>1,"DUPLICATE in A1",'
                  f'IF(COUNTIF({DBK},A{R})>1,"DUPLICATE in DB",'
                  f'IF(AND(B{R}="YES",C{R}="NO"),"MISSING in A1",'
                  f'IF(AND(B{R}="NO",C{R}="YES"),"EXTRA in A1",'
                  f'IF(J{R}<>"","VALUE MISMATCH","OK / Matched")))))'))
    ws.cell(R,12,(f'=IF(LEFT(K{R},9)="DUPLICATE","Key ซ้ำ (A1="&COUNTIF({A1K},A{R})&" / DB="&COUNTIF({DBK},A{R})&" รายการ) — ตรวจสอบ/ลบรายการซ้ำก่อน Reconcile",'
                  f'IF(K{R}="MISSING in A1","มีใน Database แต่ขาดใน A1_System — บันทึกเพิ่มที่ A1_System",'
                  f'IF(K{R}="EXTRA in A1","พบใน A1_System แต่ไม่มีใน Database — ตรวจ/ลบรายการเกินที่ A1_System",'
                  f'IF(K{R}="VALUE MISMATCH","ค่าไม่ตรง ["&J{R}&"] DB(CCY="&D{R}&",Amt="&TEXT(F{R},"#,##0")&",Dt="&H{R}&") vs A1(CCY="&E{R}&",Amt="&TEXT(G{R},"#,##0")&",Dt="&I{R}&") — แก้ที่ A1_System",""))))'))
    for c in range(2,12): st(ws.cell(R,c),bd=True,h="center" if c in(2,3) else ("right" if c in(6,7) else "left"))
    ws.cell(R,6).number_format="#,##0"; ws.cell(R,7).number_format="#,##0"
    st(ws.cell(R,11),bd=True,b=True,h="center"); st(ws.cell(R,12),bd=True,wrap=True,color=REDT)
    r+=1
last=r-1
for col,w in zip("ABCDEFGHIJKL",[12,8,8,9,9,13,13,12,12,14,16,52]): ws.column_dimensions[col].width=w
rng=f"A6:L{last}"
ws.conditional_formatting.add(rng,FormulaRule(formula=['$K6="OK / Matched"'],fill=PatternFill("solid",fgColor=GREEN),font=Font(name=F,color=GREENT)))
ws.conditional_formatting.add(rng,FormulaRule(formula=['OR($K6="EXTRA in A1",$K6="MISSING in A1")'],fill=PatternFill("solid",fgColor=RED),font=Font(name=F,color=REDT,bold=True)))
ws.conditional_formatting.add(rng,FormulaRule(formula=['$K6="VALUE MISMATCH"'],fill=PatternFill("solid",fgColor=ORANGE),font=Font(name=F,color=ORANGET,bold=True)))
ws.conditional_formatting.add(rng,FormulaRule(formula=['LEFT($K6,9)="DUPLICATE"'],fill=PatternFill("solid",fgColor=RED),font=Font(name=F,color=REDT,bold=True)))

# ============================================================ EXPORT 1B (write-back preview)
ws=wb.create_sheet("Export_1B")
title(ws,"EXPORT 1B  —  ตัวอย่างผลลัพธ์ที่จะเขียนกลับไฟล์ 1_B009  (Set D1-1)",
      "รูปแบบตรงกับไฟล์ 1_Bxxx จริง: [Key , ERROR Message]  |  ดึงเฉพาะรายการที่เป็น Error จาก Recon_Detail (แถวว่าง = ไม่มี Error, engine จริงจะ export เฉพาะแถวที่ไม่ว่าง)",3)
ws.sheet_view.showGridLines=False
hdr(ws,4,["Fi Arrangement Number (Key)","ERROR Message"],start=2)
r=5
for i in range(6,12):
    ws.cell(r,2,f'=IF(OR(Recon_Detail!K{i}="OK / Matched",Recon_Detail!K{i}=""),"",Recon_Detail!A{i})')
    ws.cell(r,3,f'=IF(OR(Recon_Detail!K{i}="OK / Matched",Recon_Detail!K{i}=""),"",Recon_Detail!L{i})')
    st(ws.cell(r,2),bd=True,b=True); st(ws.cell(r,3),bd=True,wrap=True,color=REDT)
    r+=1
elast=r-1
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=26; ws.column_dimensions["C"].width=95
ws.conditional_formatting.add(f"B5:C{elast}",FormulaRule(formula=['$B5<>""'],fill=PatternFill("solid",fgColor=RED)))
r+=1
ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
st(ws.cell(r,2,"จำนวนรายการที่จะ Export:"),b=True)
ws.cell(r,4,f'=COUNTIF(B5:B{elast},"?*")'); st(ws.cell(r,4),b=True,h="center",fill=YELLOW,bd=True)

# ============================================================ MAP CONFIG
ws=wb.create_sheet("Map_Config")
title(ws,"MAP CONFIG  —  แผนผัง Reconcile (จาก 1_Flow_Map) + Key/Date ของ Source#1",
      "สายการตรวจ Source#1→#4, Pattern, Key Mapping & คอลัมน์วันที่ (จาก 3_DataSource)",13)
ws.sheet_view.showGridLines=False
hdr(ws,4,["Product","Cat","Set","Source#1","Source#2","Source#3","Source#4","Trans.Map","Pattern","Key (S#1)","Date Filter (S#1)","Database","A1_System / Output"])
ws.freeze_panes="A5"
r=5
for p,cat,no,chain,tm,pat,mk,cmp,cond,bc in sets:
    s=chain+[""]*(4-len(chain))
    db=[x for x in chain if x[0] in "234"]; a1=[x for x in chain if x.startswith("1_A")]; out=[x for x in chain if x.startswith("1_B")]
    row=[p,cat,no,s[0],s[1],s[2],s[3],tm,pat,keymap.get(chain[0],""),datemap.get(chain[0],""),", ".join(db),(", ".join(a1+out) or "-")]
    for i,v in enumerate(row): st(ws.cell(r,i+1,v),bd=True,sz=9,h="center" if i in(0,1,2,7,8) else "left")
    st(ws.cell(r,1),b=True,fill=GREY,h="center"); r+=1
for i,w in enumerate([8,6,5,10,10,10,10,12,15,22,17,15,22]): ws.column_dimensions[get_column_letter(i+1)].width=w

# ============================================================ SOURCE CONFIG
ws=wb.create_sheet("Source_Config")
title(ws,"SOURCE CONFIG  —  เฉพาะ Field ที่จำเป็นของแต่ละไฟล์ (จาก 3_DataSource, Flag=Y)",
      "Date Filter (กรองวันที่) + Key Mapping (คีย์เชื่อม) + จำนวนคอลัมน์ Select | ไฟล์ 1_Bxxx = ไฟล์ผลลัพธ์ ERROR (key+message)",6)
ws.sheet_view.showGridLines=False
selcount={'3_O001':15,'1_A007':6,'1_A005':9,'1_A001':9,'1_A002':6,'3_O003':6,'3_O002':13,'1_A008':6,'4_C001':0,
'4_C002':2,'4_C003':2,'2_D001':15,'1_A009':12,'2_D002':52,'1_A003':52,'1_A004':6,'1_A006':5,'2_D004':5,'2_D005':5,
'1_B001':2,'1_B002':2,'1_B003':2,'1_B004':2,'1_B005':2,'1_B006':2,'1_B007':2,'1_B008':2,'1_B009':2,'2_D003':8}
group=[('1_A','A1_System (main)'),('1_B','A1_System (ERROR output)'),('2_D','Database D'),('3_O','Database O'),('4_C','Database C')]
role=[('1_A','ระบบที่ตรวจ (เกิน/ขาด)'),('1_B','ไฟล์ผลลัพธ์ ERROR'),('2_D','ฐานข้อมูล D'),('3_O','ฐานข้อมูล O'),('4_C','ฐานข้อมูล C')]
hdr(ws,4,["Group","Source / File","Date Filter Column","Key Mapping Column(s)","# Cols Selected","Role"])
ws.freeze_panes="A5"
order=['2_D001','2_D002','2_D003','2_D004','2_D005','3_O001','3_O002','3_O003','4_C001','4_C002','4_C003',
'1_A001','1_A002','1_A003','1_A004','1_A005','1_A006','1_A007','1_A008','1_A009',
'1_B001','1_B002','1_B003','1_B004','1_B005','1_B006','1_B007','1_B008','1_B009']
r=5
for s in order:
    g=next((v for k,v in group if s.startswith(k)),""); ro=next((v for k,v in role if s.startswith(k)),"")
    row=[g,s,datemap.get(s,""),keymap.get(s,""),selcount.get(s,""),ro]
    for i,v in enumerate(row): st(ws.cell(r,i+1,v),bd=True,sz=9,h="center" if i==4 else "left")
    if s.startswith("1_B"): st(ws.cell(r,2),b=True,fill=YELLOW)
    r+=1
for i,w in enumerate([20,16,24,30,15,26]): ws.column_dimensions[get_column_letter(i+1)].width=w

# ============================================================ FIELD FORMAT (normalization map)
ws=wb.create_sheet("Field_Format")
title(ws,"FIELD FORMAT  —  แผนที่ Normalize ค่า Date & Number ต่อ (Source, Field)",
      "Tool มาตรฐาน: Date = dd/mm/yyyy (internal canonical = ISO yyyy-mm-dd) | Number = ตัวเลขล้วน (ABS/แยกเครื่องหมาย) | Engine ต้อง Normalize ค่าให้เป็น canonical ก่อนเทียบทุกครั้ง  |  ช่องเหลือง = config โปรดยืนยัน, ⚠️ = ค่า default",8)
ws.sheet_view.showGridLines=False
hdr(ws,4,["Source","Field","Logical Role","Source Format (raw) ⚠️","Normalization Rule ⚠️","Compare As (canonical)","Example (raw → canonical)","Notes"])
ws.freeze_panes="A5"
# concrete date formats for demo sources; others default to confirm
date_fmt_override={('2_D001','TRADE_DATE'):('NUMBER : yyyymmdd','PARSE_DATE → ISO','20260530 → 2026-05-30','live demo'),
                   ('1_A009','Data Set Date'):('TEXT : dd/mm/yyyy','PARSE_DATE → ISO','"30/05/2026" → 2026-05-30','live demo')}
rows=[]
for s in order:
    df=datemap.get(s,'(none)')
    if df and df!='(none)':
        ov=date_fmt_override.get((s,df))
        if ov: rows.append((s,df,"Date",ov[0],ov[1],"ISO yyyy-mm-dd (แสดง dd/mm/yyyy)",ov[2],ov[3]))
        else:  rows.append((s,df,"Date","DATE ⚠️","PARSE_DATE → ISO","ISO yyyy-mm-dd (แสดง dd/mm/yyyy)","— → 2026-05-30","ยืนยัน format จริงของไฟล์"))
# number / sign rows (representative)
num_rows=[
 ("2_D001","RCV_INIT_PRINCIPAL","Number","NUMBER (อาจติดลบ)","ABS","ตัวเลข ≥ 0","-1,000,000 → 1,000,000","ตัด sign ก่อนเทียบ"),
 ("1_A009","Initial Buy Amount","Number","NUMBER (อาจติดลบ)","ABS","ตัวเลข ≥ 0","-1,000,000 → 1,000,000","ตัด sign ก่อนเทียบ"),
 ("4_C001","Amount","Number","NUMBER","SIGN_BY(\"DR/CR\", DR=+1, CR=-1)","ตัวเลขมีเครื่องหมาย","100 (CR) → -100","sign มาจากคอลัมน์ DR/CR"),
 ("2_D003","Amount","Number","NUMBER","SIGN_BY(\"Pay/Rcv\", Rcv=+1, Pay=-1)","ตัวเลขมีเครื่องหมาย","100 (Pay) → -100","sign มาจาก Pay/Rcv"),
 ("1_A008","Transaction Amount","Number","NUMBER","SIGN_BY(\"Receive Payment Type\")","ตัวเลขมีเครื่องหมาย","ตาม type","ระบุ map ค่า type → sign"),
 ("(any)","Amount (paren)","Number","TEXT : (123)","PAREN_NEG","ตัวเลขมีเครื่องหมาย","\"(100)\" → -100","วงเล็บ = ค่าลบ"),
 ("(any)","Amount (str)","Number","TEXT : \"1,000.00\"","STRIP(\",$ \") + DEC(2)","ตัวเลข","\"1,000.00\" → 1000","ตัดคั่นหลักพัน/สัญลักษณ์"),
]
allrows=rows+num_rows
r=5
for s,f,role_,sf,nr,ca,ex,nt in allrows:
    vals=[s,f,role_,sf,nr,ca,ex,nt]
    for i,v in enumerate(vals): st(ws.cell(r,i+1,v),bd=True,sz=9,h="center" if i in(2,) else "left")
    ws.cell(r,4).fill=PatternFill("solid",fgColor=INPUT); ws.cell(r,5).fill=PatternFill("solid",fgColor=INPUT)  # editable config
    if role_=="Date": st(ws.cell(r,3),color="2E75B6",b=True)
    else: st(ws.cell(r,3),color="538135",b=True)
    r+=1
flast=r-1
for i,w in enumerate([12,24,11,22,30,26,26,26]): ws.column_dimensions[get_column_letter(i+1)].width=w
r+=1
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)
st(ws.cell(r,1,'Date Source Format tokens:  DATE | NUMBER:yyyymmdd | NUMBER:ddmmyyyy | TEXT:dd/mm/yyyy | TEXT:mm/dd/yyyy | TEXT:yyyy-mm-dd | TEXT:dd-mmm-yy | TEXT:dd-mmm-yyyy'),
   b=True,sz=9,color=NAVY,fill=YELLOW,wrap=True); ws.row_dimensions[r].height=26; r+=1
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)
st(ws.cell(r,1,'Number Normalization tokens:  RAW | ABS | NEG | SIGN_BY("field",map) | PAREN_NEG | TRAIL_NEG | STRIP("chars") | DEC(n) | ROUND(n)        '
              'หมายเหตุ: yy → ปี (pivot 20yy ถ้า yy<50) , mmm = ชื่อเดือนอังกฤษ (Jan..Dec)'),
   b=True,sz=9,color=NAVY,fill=YELLOW,wrap=True); ws.row_dimensions[r].height=30

# ============================================================ OPEN QUESTIONS
ws=wb.create_sheet("Open_Questions")
title(ws,"OPEN QUESTIONS  —  สิ่งที่ยืนยันแล้ว & ที่ยังรอข้อมูล",
      "ข้อ 1-4 = ยืนยันแล้วจาก feedback | ข้อ 5+ = ยังรอเพื่อทำ Production",4)
ws.sheet_view.showGridLines=False
for c,w in zip("ABCD",[3,5,52,42]): ws.column_dimensions[c].width=w
hdr(ws,4,["","#","ประเด็น","สรุป / สถานะ"],start=1)
qs=[("1","Key mapping ข้ามไฟล์","✓ คีย์ธุรกิจเดียวกัน + บาง Source ต้อง Prep ก่อน (เช่น KT…C123→C123) → ใช้ Key Prep DSL",GREEN),
("2","เทียบค่า","✓ CCY+Amount+Date เป็นหลัก + บาง Set มีเงื่อนไขเพิ่ม (Product/Transaction type) → 3 Patterns",GREEN),
("3","Cardinality 'n'","✓ จับคู่ด้วย Match Key (ไม่รวมยอด); key มีหลายแบบ Ref / Ref+CCY+Amount / Ref+Type+CCY+Amount",GREEN),
("4","ขอบเขต","✓ ครบ 21 Sets + รองรับ chain 3-4 ขั้น (Pattern 3 Map+Back-check)",GREEN),
("5","Pattern/Compare ต่อ Set","⧖ ค่าใน Recon_Rules เป็น default โปรดยืนยัน/แก้ (โดยเฉพาะ Set ที่ตั้ง Pattern 2/3)",YELLOW),
("6","Key Prep rule จริงต่อ Source","⧖ ต้องการรูปแบบคีย์ดิบจริงของแต่ละไฟล์ (โดยเฉพาะที่ต้อง Prep) เพื่อเขียน DSL ให้ตรง",YELLOW),
("7","Pattern 3 back-check","⧖ ระบุว่าย้อนไปตรวจ field ใดที่ Main Source (เช่น D2-5→2_D001, O1-3→3_O003, C1-2→4_C003)",YELLOW),
("8","Date/Number Format ต่างกัน","✓ มีชีท Field_Format จัดการ Normalize (Date→ISO/dd/mm/yyyy, Number→ABS/SIGN_BY) — แต่ต้องยืนยัน format ดิบจริงต่อ field",GREEN),
("9","ข้อมูลตัวอย่างจริง","⧖ ขอไฟล์จริง 1-2 ไฟล์ เพื่อ map คอลัมน์ + เขียน engine (Excel/Power Query/Python) ให้รันครบทุก Set",YELLOW)]
r=5
for n,q,a,col in qs:
    st(ws.cell(r,2,n),b=True,bd=True,h="center"); st(ws.cell(r,3,q),wrap=True,bd=True,sz=9)
    st(ws.cell(r,4,a),wrap=True,bd=True,sz=9,fill=col); ws.row_dimensions[r].height=42; r+=1

order_sheets=["README","Control_Panel","Recon_Summary","Recon_Rules","Mapping_Steps","Recon_Detail","Export_1B","Sample_Data","Field_Format","Map_Config","Source_Config","Open_Questions"]
wb._sheets.sort(key=lambda s: order_sheets.index(s.title))
wb.calculation.fullCalcOnLoad=True
wb.save("Part_Reconcile_Tool_Prototype.xlsx")
print("saved",len(sets),"sets")
